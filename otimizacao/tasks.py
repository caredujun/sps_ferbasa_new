#import xlwt
import io
from django.core.mail import EmailMessage
from celery import shared_task
from django.core.mail import send_mail
from django.http import FileResponse, HttpResponse


from produtos.models import TbProdutoMercadoPreco, TbProdutoMercadoPrecoDaugther
from tabelas.models import TbEquacaoAjustePreco
from .models import *
from django.db import connection

locale.setlocale(locale.LC_ALL, 'pt_BR.utf8')  # Estou usando esse pois Heroku não aceita pt_BR
locale.setlocale(locale.LC_ALL, 'pt_BR.utf8')  # Estou usando esse pois Heroku não aceita pt_BR

@shared_task
def exportar_excel_tbprodutomercadofluxo_celery(lista_id, user_mail):

    #Comentamos pois iremos enviar o arquivo ao usuário e não mostrar para download
    #response = HttpResponse(content_type='application/ms-excel')
    #response['Content-Disposition'] = 'attachment; filename="Otimização Produtos Mercados Fluxos.xls"'

    import openpyxl

    # Call a Workbook() function of openpyxl
    # to create a new blank Workbook object
    wb = openpyxl.Workbook()

    # Get workbook active sheet
    # from the active attribute
    ws = wb.active

    #Old
    #wb = xlwt.Workbook(encoding='utf-8')

    # Lançando valores das filhas
    #New
    # One can change the name of the title
    ws.title = "resultados"
    #Old
    #ws = wb.add_sheet('resultados')

    filhas = TbProdutoMercadoFluxoDaugther.objects.filter(mae_id__in=lista_id).order_by('id')

    # Sheet header, first row
    row_num = 1

    #Old
    #font_style = xlwt.XFStyle()
    #font_style.font.bold = True

    # Vamos ver o tipo de período
    if TbCenarios.objects.get(cen_ativo=True).cen_tipo == 'Anual':
        tipo_periodo = 'Ano'
    elif TbCenarios.objects.get(cen_ativo=True).cen_tipo == 'Trimestral':
        tipo_periodo = 'Ano/Trimestre'
    else:
        tipo_periodo = 'Ano/Mês'


    # Vamos pegar o cenário ativo (id e nome)
    cen_ativo = TbCenarios.objects.get(cen_ativo=True).id
    cen_ativo_nome = TbCenarios.objects.get(cen_ativo=True).cen_nome

    #New
    ws.cell(row=row_num, column=1).value = 'RESULTADOS OTIMIZAÇÃO: PRODUTO-MERCADO-FLUXO DE PRODUÇÃO / CENÁRIO ' + str(cen_ativo) + ' / ' + cen_ativo_nome
    #Old
    #ws.write(row_num, 0, 'RESULTADOS OTIMIZAÇÃO: PRODUTO-MERCADO-FLUXO DE PRODUÇÃO / CENÁRIO ' + str(cen_ativo) + ' / ' + cen_ativo_nome, font_style)

    columns = [tipo_periodo, 'Produto', 'Mercado', 'Preço Comercial', 'Unidade', 'Fluxo de Produção', 'Fluxo Running', 'Equip. Running', 'Vendas', 'Margen Contrib.', 'Margen Horária', 'Equip. Margem Horária', 'Preço', 'Custo Variável', 'Parcela Itens', 'Parcela Inbound', 'Parcela Outbound',
               'Parcela Manutenção', 'Total Margem']

    row_num += 1
    for col_num in range(len(columns)):
        #new
        ws.cell(row=row_num, column=col_num+1).value = columns[col_num]
        #Old
        #ws.write(row_num, col_num, columns[col_num], font_style)

    #Old
    # Sheet body, remaining rows
    #font_style = xlwt.XFStyle()

    rows = filhas.values_list('dau_order', 'dau_valor_1', 'dau_valor_7', 'dau_valor_2', 'dau_valor_3', 'dau_valor_4', 'dau_valor_5', 'dau_valor_8', 'dau_valor_9', 'dau_valor_6', 'dau_valor_10', 'mae_id').order_by('mae_id', 'dau_order')

    # Vamos incluir os nomes de algumas colunas nos valores da lista new_rows
    new_rows = []
    linha_num = 0
    for row in rows:
        id_mae = row[11]

        # Vamos pegar os id e nomes
        id_produto = TbProdutoMercadoFluxo.objects.get(id=id_mae).pro_mer_flu_produto_id
        nome_produto = TbProdutos.objects.get(id=id_produto).pro_codigo

        id_mercado = TbProdutoMercadoFluxo.objects.get(id=id_mae).pro_mer_flu_mercado_id
        nome_mercado = TbMercado.objects.get(id=id_mercado).mer_nome

        # Vamos pegar o preço comercial. Está na tabela TbProdutoMercadoPrecoDaugther
        # Temos que pegar o id na tabela mãe TbProdutoMercadoPreco
        id_produto_mercado = TbProdutoMercadoPreco.objects.get(pro_mer_pre_produto_id=id_produto, pro_mer_pre_mercado_id=id_mercado).id
        if TbProdutoMercadoPrecoDaugther.objects.filter(mae_id=id_produto_mercado, dau_order=row[0]).exists():
            produto_mercado_preco = TbProdutoMercadoPrecoDaugther.objects.get(mae_id=id_produto_mercado, dau_order=row[0]).dau_valor_3
        else:
            produto_mercado_preco = 0

        # Vamos pegar a unidade do preço comercial
        if TbProdutoMercadoPreco.objects.get(pro_mer_pre_produto_id=id_produto, pro_mer_pre_mercado_id=id_mercado).pro_mer_pre_equacao_id:
            # Foi informado equação de preço. Vamos pegar a unidade do preço comercial na tabela TbEquacaoAjustePreco
            unidade_preco_comercial = TbEquacaoAjustePreco.objects.get(id=TbProdutoMercadoPreco.objects.get(pro_mer_pre_produto_id=id_produto, pro_mer_pre_mercado_id=id_mercado).pro_mer_pre_equacao_id).equ_aju_pre_descricao
        else:
            # Não foi informado equação. Vamos pegar a unidade de preço comercial na tabelaTbProdutoMercadoPreco
            unidade_preco_comercial = TbProdutoMercadoPreco.objects.get(pro_mer_pre_produto_id=id_produto, pro_mer_pre_mercado_id=id_mercado).pro_mer_pre_moeda

        id_fluxo = TbProdutoMercadoFluxo.objects.get(id=id_mae).pro_mer_flu_fluxo_producao_id
        nome_fluxo = TbFluxoProducao.objects.get(id=id_fluxo).flu_pro_descricao

        # Margem Horária
        cursor = connection.cursor()
        # Expressão SQL
        sql = "call public.gargalo_produtividade_fluxo_producao_order_valor(" + str(id_fluxo) + ", " + str(row[0]) + ", 0)"
        cursor.execute(sql)
        retorno = cursor.fetchone()[0]
        cursor.close()
        # Retorna o valor da produtividade no equipamento selecionado para ser considerado na margem de contribuição horária.
        # Se não foi selecionado nenhum equipamento, retorna ZERO.
        # Se multiplicarmos pela margem de contribuição bruta (dau_valor_3=row[4]) temos a margem horária.
        if retorno > 0:
            margem_horaria = round(retorno * row[4], 2)
        else:
            margem_horaria = 'ND'

        # Equipamento Margem Horária
        cursor = connection.cursor()
        # Vamos montar uma expressão SQL para mostrar o equipamento selecionado para definir a margem horária
        # Expressão SQL
        sql = "call public.gargalo_produtividade_fluxo_producao_order_equipamento(" + str(id_fluxo) + ", " + str(row[0]) + ", '')"
        cursor.execute(sql)
        equipamento_margem_horaria = cursor.fetchone()[0]
        cursor.close()

        # Inserindo na lista new_rows
        list_aux = list(rows[linha_num])
        list_aux.insert(1, nome_produto)
        list_aux.insert(2, nome_mercado)
        list_aux.insert(3, produto_mercado_preco)
        list_aux.insert(4, unidade_preco_comercial)
        list_aux.insert(5, nome_fluxo)
        list_aux.insert(10, margem_horaria)
        list_aux.insert(11, equipamento_margem_horaria)
        # Vamos remover o mae_id da lista. Está na posição 18 após as inclusões dos nomes.
        del (list_aux[18])
        # Vamos incluir o total da margem de contribuição
        total_margem = list_aux[8] * list_aux[9]
        list_aux.insert(18, total_margem)

        # Vamos agora alterar o dau_order para mostrar o período no formato (ano, ano/mês ou ano/trimestre)
        # Temos que pegar o inicio do cenário
        inicio_periodo = TbCenarios.objects.get(cen_ativo=True).cen_inicio

        if TbCenarios.objects.get(cen_ativo=True).cen_tipo == 'Anual':
            # Temos que considerar o inicio-periodo somente os 4 digitos iniciais
            inicio_periodo = inicio_periodo[:4]
            periodo_str = str(int(inicio_periodo) - 1 + list_aux[0])

        if TbCenarios.objects.get(cen_ativo=True).cen_tipo == 'Mensal':
            year = int(inicio_periodo[:4])
            month = int(inicio_periodo[-2:])

            for j in range(list_aux[0] - 1):
                if month + 1 > 12:
                    month = 1
                    year = year + 1
                else:
                    month = month + 1
            if month < 10:
                periodo_str = str(year) + '/0' + str(month)
            else:
                periodo_str = str(year) + '/' + str(month)

        if TbCenarios.objects.get(cen_ativo=True).cen_tipo == 'Trimestral':
            year = int(inicio_periodo[:4])
            quarter = int(inicio_periodo[-2:])
            for j in range(list_aux[0] - 1):
                if quarter + 1 > 4:
                    quarter = 1
                    year = year + 1
                else:
                    quarter = quarter + 1
            periodo_str = str(year) + '/0' + str(quarter)

        list_aux[0] = periodo_str

        new_rows.append(list_aux)

        linha_num += 1

    rows = new_rows
    for row in rows:
        row_num += 1
        for col_num in range(len(row)):
            #New
            ws.cell(row=row_num, column=col_num+1).value = row[col_num]
            #Old
            #ws.write(row_num, col_num, row[col_num], font_style)

    # Arquivo foi gerado. Vamos enviar para o usuário via e-mail
    lista_e_mail = []
    lista_e_mail.append(user_mail)

    mail = EmailMessage(subject="Arquivo Excel Sistema SPS",
                        from_email='sps.consultoria.alerta@gmail.com',
                        to=lista_e_mail,
                        body="Arquivo gerado com sucesso. Tabela Otimização - Produtos/Mercados/Fluxos")

    with io.BytesIO() as xlsx:
        wb.save(xlsx)
        mail.attach(filename="Otimização Produtos Mercados Fluxos.xlsx",
                    content=xlsx.getvalue(),
                    mimetype="application/ms-excel")
    #print('Enviando e-mail')
    mail.send(fail_silently=False)
    #print('Enviando e-mail')


# ---------------------------------------------------------------------------------------------------------------
# 🌟 NOVO (18/10/2026): análises do cenário ativo pedidas pelo Agente IA (chat). Rodam aqui, em segundo plano (no Heroku
# uma requisição de mais de 30 s derruba o sistema). O resultado (texto da resposta do chat, com o quadro) fica no
# resultado da task; o chat sonda e mostra quando terminar (etapa 'cfx_processando' em parameters/fluxo_criar_cenario.py).
#   - Análise de Resultados: otimizacao/analise_resultados.py (lê as tabelas preenchidas por Otimizar / Consolidar)
#   - Sensibilidade Custo x Produtividade: otimizacao/sensibilidade.py (o mesmo modelo do Otimizar --
#     parameters.tasks.otimizar_cenario_celery -- resolvido de novo com scipy.optimize.linprog, method='highs')
# ---------------------------------------------------------------------------------------------------------------
@shared_task(bind=True)
def analise_resultados_chat_celery(self, usuario_id, dados):
    import time as _time
    from parameters.fluxo_criar_cenario import montar_analise_otimizacao_em_segundo_plano
    comeco = _time.time()
    resposta = montar_analise_otimizacao_em_segundo_plano(usuario_id, 'analise_resultados', dados)
    print(f"[otimizacao][analise_resultados] pronta em {_time.time() - comeco:.1f}s ({len(resposta or '')} caracteres)")
    return resposta


@shared_task(bind=True)
def sensibilidade_custo_produtividade_chat_celery(self, usuario_id, dados):
    import time as _time
    from parameters.fluxo_criar_cenario import montar_analise_otimizacao_em_segundo_plano
    comeco = _time.time()
    resposta = montar_analise_otimizacao_em_segundo_plano(usuario_id, 'sensibilidade', dados)
    print(f"[otimizacao][sensibilidade] pronta em {_time.time() - comeco:.1f}s ({len(resposta or '')} caracteres)")
    return resposta


@shared_task(bind=True)
def sensibilidade_ganho_margem_chat_celery(self, usuario_id, visao, dados):
    """🌟 NOVO (20/10/2026): "Sensibilidade Ganho Margem x Produtividade" (visao 'sens_ganho_prod') e
    "Sensibilidade Ganho Margem x Limite Comercial" (visao 'sens_ganho_limite') -- otimizacao/sensibilidade_shadow.py."""
    import time as _time
    from parameters.fluxo_criar_cenario import montar_analise_otimizacao_em_segundo_plano
    comeco = _time.time()
    resposta = montar_analise_otimizacao_em_segundo_plano(usuario_id, visao, dados)
    print(f"[otimizacao][{visao}] pronta em {_time.time() - comeco:.1f}s ({len(resposta or '')} caracteres)")
    return resposta
